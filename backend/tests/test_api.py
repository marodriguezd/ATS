"""API integration: upload -> audit -> auto-fix -> export, no fixed IDs/state."""
import io

import docx


def _make_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_paragraph("Alex Example")
    doc.add_paragraph("alex@example.com | +34 600 111 222")
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("Built REST endpoints with Python.")
    doc.add_paragraph("EDUCATION")
    doc.add_paragraph("Higher Technician in DAM")
    doc.add_paragraph("TECHNICAL SKILLS")
    doc.add_paragraph("Python, SQL")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_full_workflow_txt_upload(client):
    r = client.post("/api/resumes/upload", files={"file": ("cv.txt", b"Alex Example\nalex@example.com | +34 600 111 222\n\nWORK EXPERIENCE\nBuilt APIs.\n\nEDUCATION\nDAM.\n\nTECHNICAL SKILLS\nPython.", "text/plain")})
    assert r.status_code == 200, r.text
    rid = r.json()["id"]

    audit = client.post("/api/audit/run", json={"resume_id": rid, "job_text": "We require Python and SQL."})
    assert audit.status_code == 200
    assert "methodology_note" in audit.json()["result"]

    fix = client.post("/api/audit/auto-fix", json={"resume_id": rid, "job_text": "Python role"})
    assert fix.status_code == 200
    body = fix.json()
    assert "miguadali" not in str(body).lower()
    new_id = body["new_resume_id"]

    exp = client.get(f"/api/resumes/{new_id}/export/txt")
    assert exp.status_code == 200
    assert "alex@example.com" in exp.text


def test_docx_upload_and_export_pdf(client):
    r = client.post("/api/resumes/upload", files={"file": ("cv.docx", _make_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code == 200, r.text
    rid = r.json()["id"]
    pdf = client.get(f"/api/resumes/{rid}/export/pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")


def test_missing_resume_is_404_not_substitute(client):
    r = client.get("/api/resumes/999999")
    assert r.status_code == 404
    a = client.post("/api/audit/run", json={"resume_id": 999999})
    assert a.status_code == 404


def test_invalid_uploads_rejected(client):
    r = client.post("/api/resumes/upload", files={"file": ("x.exe", b"MZ...", "application/octet-stream")})
    assert r.status_code == 400
    r2 = client.post("/api/resumes/upload", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")})
    assert r2.status_code == 400
