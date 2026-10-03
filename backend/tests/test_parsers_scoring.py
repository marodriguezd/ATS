"""Parser adversarial suite + scoring sanity."""
import io
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.core.ats_parser import ATSParser
from app.core.scorer import ATSScorer


def test_accented_spanish_and_contact():
    text = ("María López\nmaria@example.es | +34 612 345 678\n\n"
            "EXPERIENCIA LABORAL\nSoporte con Python.\n\nFORMACIÓN\nDAM.\n\n"
            "HABILIDADES\nPython, SQL.")
    parsed = ATSParser.parse_plain_text(text)
    assert parsed["contact_info"]["email"] == "maria@example.es"
    assert parsed["contact_info"]["phone"] == "+34 612 345 678"
    assert set(parsed["sections"]) >= {"experience", "education", "skills"}


def test_years_are_not_phones():
    parsed = ATSParser.parse_plain_text("EDUCATION\nDAM (2021 - 2023)\nNo contact here.")
    assert parsed["contact_info"]["phone"] is None


def test_docx_tables_and_txt_roundtrip():
    import docx
    doc = docx.Document()
    doc.add_paragraph("Alex Example")
    doc.add_paragraph("alex@example.com | +34 600 111 222")
    doc.add_paragraph("WORK EXPERIENCE")
    doc.add_paragraph("Built REST endpoints.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    buf = io.BytesIO()
    doc.save(buf)
    parsed = ATSParser.parse_docx(buf.getvalue())
    assert parsed["has_tables"] is True
    assert "alex@example.com" in parsed["raw_text"]
    assert "Python" in parsed["raw_text"]


def test_job_zones_and_weights():
    resume = ATSParser.parse_plain_text(
        "Alex\nalex@example.com | +34 600 111 222\n\nWORK EXPERIENCE\n"
        "Built Python APIs.\n\nEDUCATION\nDAM.\n\nTECHNICAL SKILLS\nPython, SQL.")
    job = "We require Python and SQL. Nice to have: Kubernetes."
    result = ATSScorer.score_all(resume, job)
    tags = {k["keyword"]: k["tag"] for k in result["keyword_details"]["keywords"]}
    assert tags.get("kubernetes") == "PREFERRED"
    assert result["weights"]["keyword_match"] == 0.40
    assert "methodology_note" in result
    # no-job scoring excludes keyword weight
    plain = ATSScorer.score_all(resume, "")
    assert plain["breakdown"]["keyword_match"] == 0
    assert "keyword_match" not in plain["weights"]


def test_bare_numbers_are_not_evidence():
    resume = ATSParser.parse_plain_text(
        "Alex\nalex@example.com\n\nWORK EXPERIENCE\nWorked in 2023 on 3 projects.\n\n"
        "EDUCATION\nDAM.\n\nTECHNICAL SKILLS\nPython.")
    result = ATSScorer.score_all(resume, "")
    assert result["impact_details"]["total_metrics_found"] == 0
